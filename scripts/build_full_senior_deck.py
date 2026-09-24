# -*- coding: utf-8 -*-
"""
scripts/build_full_senior_deck.py
構建《高三理組數學思維本(2026)》全書互動教學簡報與 OMML 核心系統：
1. 讀取全部 20 個專題章節 JS 文件
2. 批量生成全部題目之微軟 Word 原生 OMML 數學方程式
3. 注入 omml 屬性至各專題 JS 文件
4. 生成《高三理組數學思維本_2026_參考答案_OMML對照手冊.docx》與《.xml》彙編
5. 升級 omml.html 在線檢索中心，涵蓋全書 20 大專題
6. 創建 deck2026.html 主入口簡報，標題為《高三理組數學思維本(2026)》
"""

import os
import re
import json
import shutil
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from latex2mathml.converter import convert as latex2mathml_convert
from lxml import etree

BASE_DIR = r"g:\我的雲端硬碟\12_數學教材轉化\tools\math-review-deck"
DEMO_DIR = os.path.join(BASE_DIR, "demo")
ASSETS_DIR = os.path.join(BASE_DIR, "assets")

XSL_PATH = r"C:\Program Files\Microsoft Office\root\Office16\MML2OMML.XSL"
xslt = etree.parse(XSL_PATH)
transform = etree.XSLT(xslt)

CHAPTER_FILES = [
    "ch-calc-01.js",
    "ch-calc-02.js",
    "ch-calc-03.js",
    "ch-calc-04.js",
    "ch-calc-05.js",
    "ch-senior-06.js",
    "ch-senior-07.js",
    "ch-senior-08.js",
    "ch-senior-09.js",
    "ch-senior-10.js",
    "ch-senior-11.js",
    "ch-senior-12.js",
    "ch-senior-13.js",
    "ch-senior-14.js",
    "ch-senior-15.js",
    "ch-senior-16.js",
    "ch-senior-17.js",
    "ch-senior-18.js",
    "ch-senior-19.js",
    "ch-senior-20.js"
]

def clean_latex(latex_str):
    """清洗 LaTeX 字串，轉換為純淨數學公式"""
    s = latex_str.strip()
    s = re.sub(r'^\$+|\$+$', '', s).strip()
    s = re.sub(r'\\left\.', '', s)
    s = re.sub(r'\\right\.', '', s)
    s = re.sub(r'\\quad', ' ', s)
    s = re.sub(r'\\qquad', '  ', s)
    s = re.sub(r'\\text\{([^}]+)\}', r'\1', s)
    s = re.sub(r'\\mathbb\{R\}', 'R', s)
    s = re.sub(r'\\mathbb\{N\}\^\*', 'N*', s)
    s = re.sub(r'\\mathbf\{([^}]+)\}', r'\1', s)
    s = re.sub(r'\\vec\{([a-zA-Z0-9]+)\}', r'\1', s)
    s = re.sub(r'\\triangle', '△', s)
    s = re.sub(r'\\angle', '∠', s)
    s = re.sub(r'\\circ', '°', s)
    s = re.sub(r'\\parallel', '//', s)
    s = re.sub(r'\\perp', '⊥', s)
    s = re.sub(r'\\sim', '∽', s)
    s = re.sub(r'\\cong', '≌', s)
    s = re.sub(r'\\cdot', '·', s)
    s = re.sub(r'\\times', '×', s)
    s = re.sub(r'\\div', '÷', s)
    s = re.sub(r'\\pm', '±', s)
    s = re.sub(r'\\mp', '∓', s)
    s = re.sub(r'\\le', '≤', s)
    s = re.sub(r'\\ge', '≥', s)
    s = re.sub(r'\\neq', '≠', s)
    s = re.sub(r'\\approx', '≈', s)
    s = re.sub(r'\\infty', '∞', s)
    s = re.sub(r'\\in', '∈', s)
    s = re.sub(r'\\notin', '∉', s)
    s = re.sub(r'\\subset', '⊂', s)
    s = re.sub(r'\\subseteq', '⊆', s)
    s = re.sub(r'\\cup', '∪', s)
    s = re.sub(r'\\cap', '∩', s)
    s = re.sub(r'\\emptyset', '∅', s)
    s = re.sub(r'\\iff', '<=>', s)
    s = re.sub(r'\\implies', '=>', s)
    s = re.sub(r'\\to', '->', s)
    s = re.sub(r'\\overline\{([^}]+)\}', r'\1', s)
    s = re.sub(r'\\bar\{([^}]+)\}', r'\1', s)
    s = re.sub(r'\\hat\{([^}]+)\}', r'\1', s)
    s = re.sub(r'\\binom\{([^}]+)\}\{([^}]+)\}', r'C(\1,\2)', s)
    s = re.sub(r'\\begin\{pmatrix\}(.*?)\\end\{pmatrix\}', r'[\1]', s, flags=re.DOTALL)
    s = re.sub(r'\\begin\{vmatrix\}(.*?)\\end\{vmatrix\}', r'|\1|', s, flags=re.DOTALL)
    s = re.sub(r'\\\\', '; ', s)
    s = re.sub(r'&', ' ', s)
    s = re.sub(r'\\sum_\{([^}]+)\}\^\{([^}]+)\}', r'sum_{\1}^{\2}', s)
    s = re.sub(r'\\int_\{([^}]+)\}\^\{([^}]+)\}', r'int_{\1}^{\2}', s)
    return s.strip()

def latex_to_omml(latex_code):
    """將 LaTeX 數學公式轉化為微軟 Word 原生 OMML XML"""
    clean_code = clean_latex(latex_code)
    if not clean_code:
        return ""
    
    if clean_code in ['A', 'B', 'C', 'D', 'E']:
        return f'<m:oMath xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math"><m:r><m:t>{clean_code}</m:t></m:r></m:oMath>'
    
    try:
        mml = latex2mathml_convert(clean_code)
        mml_dom = etree.fromstring(mml.encode('utf-8'))
        omml_dom = transform(mml_dom)
        omml_str = etree.tostring(omml_dom, encoding='utf-8').decode('utf-8')
        omml_clean = re.sub(r'<\?xml[^>]*\?>', '', omml_str).strip()
        omml_clean = re.sub(r'xmlns:m="[^"]*"', '', omml_clean).strip()
        omml_clean = re.sub(r'xmlns:mml="[^"]*"', '', omml_clean).strip()
        omml_clean = re.sub(r'<m:oMathPara>', '', omml_clean)
        omml_clean = re.sub(r'</m:oMathPara>', '', omml_clean).strip()
        
        if not omml_clean.startswith('<m:oMath'):
            omml_clean = f'<m:oMath xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math">{omml_clean}</m:oMath>'
        else:
            omml_clean = re.sub(r'<m:oMath[^>]*>', '<m:oMath xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math">', omml_clean, count=1)
        return omml_clean
    except Exception as e:
        safe_text = clean_code.replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;')
        return f'<m:oMath xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math"><m:r><m:t>{safe_text}</m:t></m:r></m:oMath>'

def extract_deck_from_js(js_path):
    with open(js_path, 'r', encoding='utf-8') as f:
        content = f.read()
    m = re.search(r'DECK\.push\((\{[\s\S]*\})\);?\s*\}\)\(\);?', content)
    if not m:
        raise ValueError(f"Could not parse DECK.push in {js_path}")
    return json.loads(m.group(1))

def save_deck_to_js(deck_obj, js_path):
    ch_name = deck_obj.get("ch", "")
    slides_count = len(deck_obj.get("slides", []))
    header = f"/* 2026 高三理組數學思維本 — {ch_name} ({slides_count} 題) - 支援 OMML */"
    body = json.dumps(deck_obj, ensure_ascii=False, indent=2)
    content = f"{header}\n(function() {{\n  const DECK = window.DECK = window.DECK || [];\n  DECK.push({body});\n}})();\n"
    with open(js_path, 'w', encoding='utf-8') as f:
        f.write(content)

print("Starting full deck build and OMML processing...")

all_chapters_data = []
total_slides = 0

for fname in CHAPTER_FILES:
    js_path = os.path.join(DEMO_DIR, fname)
    if not os.path.exists(js_path):
        print(f"Warning: {js_path} does not exist, skipping.")
        continue
    
    deck = extract_deck_from_js(js_path)
    slides = deck.get("slides", [])
    
    # Process OMML for each slide
    for slide in slides:
        total_slides += 1
        sol = slide.get("solution", {})
        ans = sol.get("ans", "")
        # If OMML is missing or empty, generate it
        if "omml" not in sol or not sol["omml"]:
            omml_code = latex_to_omml(ans)
            sol["omml"] = omml_code
        slide["solution"] = sol
    
    deck["slides"] = slides
    save_deck_to_js(deck, js_path)
    # also save to assets
    asset_path = os.path.join(ASSETS_DIR, fname)
    save_deck_to_js(deck, asset_path)
    
    all_chapters_data.append(deck)
    print(f"  Processed {fname}: {deck.get('title')} ({len(slides)} slides)")

print(f"\nAll {len(all_chapters_data)} chapters processed, total {total_slides} slides.")

# ==============================================================================
# 生成微軟 Word 原生 OMML 對照手冊 (.docx)
# ==============================================================================
print("\nGenerating DOCX manual: 高三理組數學思維本_2026_參考答案_OMML對照手冊.docx ...")
doc = docx.Document()

# 頁邊距
sections = doc.sections
for sec in sections:
    sec.top_margin = Inches(0.8)
    sec.bottom_margin = Inches(0.8)
    sec.left_margin = Inches(0.8)
    sec.right_margin = Inches(0.8)

# 標題
p_title = doc.add_paragraph()
r_title = p_title.add_run("高三理組數學思維本 (2026)\n全書參考答案與 OMML 數學方程式對照手冊")
r_title.bold = True
r_title.font.size = Pt(20)
r_title.font.name = "微軟正黑體"
r_title.font.color.rgb = RGBColor(30, 58, 138)
p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER

p_sub = doc.add_paragraph()
r_sub = p_sub.add_run(f"全書 20 大專題 · 總計 {total_slides} 道核心題目與題組 · 微軟 Word 原生 OMML 方程式完整收錄")
r_sub.font.size = Pt(11)
r_sub.font.name = "微軟正黑體"
r_sub.font.color.rgb = RGBColor(100, 116, 139)
p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER

doc.add_paragraph("本手冊收錄《高三理組數學思維本(2026)》全書二十大專題全部題目之標準參考答案，所有方程式均以微軟 Office Math Markup Language (OMML) 格式生成。支援在 Microsoft Word 中原生編輯、排版、縮放與精確列印。")

# 表格列出所有專題與題目
for ch_idx, ch in enumerate(all_chapters_data, 1):
    doc.add_heading(f"專題 {ch_idx}：{ch.get('title')}", level=1)
    
    slides = ch.get("slides", [])
    table = doc.add_table(rows=1, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = "題號"
    hdr_cells[1].text = "核心考點主題"
    hdr_cells[2].text = "參考答案 (LaTeX / 文字)"
    hdr_cells[3].text = "微軟原生 OMML 方程式"
    
    # Header format
    for cell in hdr_cells:
        cell.paragraphs[0].runs[0].font.bold = True
        cell.paragraphs[0].runs[0].font.size = Pt(10)
        shading = parse_xml(r'<w:shd {} w:fill="2563EB"/>'.format(nsdecls('w')))
        cell._tc.get_or_add_tcPr().append(shading)
        cell.paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
    
    for s_idx, s in enumerate(slides, 1):
        row = table.add_row()
        cells = row.cells
        cells[0].text = s.get("qNum", f"第{s_idx}題")
        cells[1].text = s.get("topic", "")
        
        sol = s.get("solution", {})
        ans = sol.get("ans", "")
        cells[2].text = ans
        
        # Insert OMML
        omml_xml = sol.get("omml", "")
        p_omml = cells[3].paragraphs[0]
        if omml_xml:
            try:
                omml_elem = parse_xml(omml_xml.encode('utf-8'))
                p_omml._p.append(omml_elem)
            except Exception as e:
                p_omml.text = ans
        else:
            p_omml.text = ans
        
        # cell widths
        cells[0].width = Inches(1.1)
        cells[1].width = Inches(2.0)
        cells[2].width = Inches(1.9)
        cells[3].width = Inches(2.0)

docx_out_demo = os.path.join(DEMO_DIR, "高三理組數學思維本_2026_參考答案_OMML對照手冊.docx")
docx_out_assets = os.path.join(ASSETS_DIR, "高三理組數學思維本_2026_參考答案_OMML對照手冊.docx")
doc.save(docx_out_demo)
shutil.copyfile(docx_out_demo, docx_out_assets)
print(f"Saved DOCX manual to:\n  - {docx_out_demo}\n  - {docx_out_assets}")

# ==============================================================================
# 生成 OMML XML 彙編文件 (.xml)
# ==============================================================================
print("\nGenerating XML compilation: 高三理組數學思維本_2026_參考答案_OMML彙編.xml ...")
xml_lines = [
    '<?xml version="1.0" encoding="utf-8"?>',
    '<MathDeckOMMLCompilation xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math" title="高三理組數學思維本 (2026) 參考答案 OMML 彙編" year="2026">'
]

for ch_idx, ch in enumerate(all_chapters_data, 1):
    ch_title = ch.get("title", "").replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
    xml_lines.append(f'  <Chapter id="{ch_idx}" title="{ch_title}">')
    for s_idx, s in enumerate(ch.get("slides", []), 1):
        q_num = s.get("qNum", f"第{s_idx}題").replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
        topic = s.get("topic", "").replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
        sol = s.get("solution", {})
        ans = sol.get("ans", "").replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
        omml = sol.get("omml", "")
        xml_lines.append(f'    <Item index="{s_idx}" qNum="{q_num}">')
        xml_lines.append(f'      <Topic>{topic}</Topic>')
        xml_lines.append(f'      <RawAnswer>{ans}</RawAnswer>')
        xml_lines.append(f'      <OMML>')
        xml_lines.append(f'        {omml}')
        xml_lines.append(f'      </OMML>')
        xml_lines.append(f'    </Item>')
    xml_lines.append('  </Chapter>')

xml_lines.append('</MathDeckOMMLCompilation>')
xml_content = "\n".join(xml_lines)

xml_out_demo = os.path.join(DEMO_DIR, "高三理組數學思維本_2026_參考答案_OMML彙編.xml")
xml_out_assets = os.path.join(ASSETS_DIR, "高三理組數學思維本_2026_參考答案_OMML彙編.xml")
with open(xml_out_demo, "w", encoding="utf-8") as f:
    f.write(xml_content)
shutil.copyfile(xml_out_demo, xml_out_assets)
print(f"Saved XML compilation to:\n  - {xml_out_demo}\n  - {xml_out_assets}")

# ==============================================================================
# 更新 omml.html 在線檢索與預覽庫
# ==============================================================================
print("\nUpdating OMML Online Portal: omml.html ...")
omml_portal_html = f"""<!DOCTYPE html>
<html lang="zh-TW">
<head>
  <meta charset="UTF-8" />
  <meta name="viewport" content="width=device-width, initial-scale=1.0" />
  <title>高三理組數學思維本 (2026) · OMML 數學方程式在線檢索庫</title>
  <meta name="description" content="高三理組數學思維本微軟 Word 原生 OMML 數學方程式檢索、預覽與一鍵複製中心" />
  <style>
    :root {{
      --bg: #0f172a;
      --card-bg: #1e293b;
      --border: #334155;
      --text: #f8fafc;
      --sub: #94a3b8;
      --accent: #38bdf8;
      --purple: #a855f7;
      --green: #22c55e;
    }}
    * {{ box-sizing: border-box; margin: 0; padding: 0; }}
    body {{
      font-family: -apple-system, BlinkMacSystemFont, "Segoe UI", Roboto, "Helvetica Neue", Arial, sans-serif;
      background: var(--bg);
      color: var(--text);
      line-height: 1.6;
      padding: 24px;
    }}
    header {{
      max-width: 1200px;
      margin: 0 auto 30px auto;
      text-align: center;
      padding-bottom: 20px;
      border-bottom: 1px solid var(--border);
    }}
    h1 {{
      font-size: 2.2rem;
      margin-bottom: 10px;
      background: linear-gradient(135deg, #38bdf8, #818cf8, #c084fc);
      -webkit-background-clip: text;
      -webkit-text-fill-color: transparent;
    }}
    p.lead {{
      color: var(--sub);
      font-size: 1.05rem;
      margin-bottom: 16px;
    }}
    .actions {{
      display: flex;
      justify-content: center;
      gap: 12px;
      flex-wrap: wrap;
      margin-top: 15px;
    }}
    .btn {{
      display: inline-flex;
      align-items: center;
      gap: 6px;
      padding: 10px 20px;
      border-radius: 8px;
      font-size: 0.95rem;
      font-weight: 600;
      text-decoration: none;
      transition: all 0.2s;
      cursor: pointer;
    }}
    .btn-primary {{
      background: #2563eb;
      color: #fff;
      border: 1px solid #3b82f6;
    }}
    .btn-primary:hover {{ background: #1d4ed8; }}
    .btn-sec {{
      background: rgba(255, 255, 255, 0.08);
      color: #cbd5e1;
      border: 1px solid var(--border);
    }}
    .btn-sec:hover {{ background: rgba(255, 255, 255, 0.15); color: #fff; }}
    .filters {{
      max-width: 1200px;
      margin: 0 auto 24px auto;
      display: flex;
      gap: 12px;
      flex-wrap: wrap;
    }}
    .search-box {{
      flex: 1;
      min-width: 250px;
      padding: 12px 16px;
      background: var(--card-bg);
      border: 1px solid var(--border);
      border-radius: 8px;
      color: #fff;
      font-size: 1rem;
    }}
    .select-box {{
      padding: 12px 16px;
      background: var(--card-bg);
      border: 1px solid var(--border);
      border-radius: 8px;
      color: #fff;
      font-size: 1rem;
    }}
    .container {{
      max-width: 1200px;
      margin: 0 auto;
      display: grid;
      grid-template-columns: repeat(auto-fill, minmax(360px, 1fr));
      gap: 20px;
    }}
    .card {{
      background: var(--card-bg);
      border: 1px solid var(--border);
      border-radius: 12px;
      padding: 20px;
      display: flex;
      flex-direction: column;
      justify-content: space-between;
      transition: transform 0.2s, border-color 0.2s;
    }}
    .card:hover {{
      transform: translateY(-2px);
      border-color: var(--accent);
    }}
    .card-top {{
      margin-bottom: 12px;
    }}
    .card-tag {{
      display: inline-block;
      font-size: 0.75rem;
      padding: 3px 8px;
      border-radius: 4px;
      background: rgba(56, 189, 248, 0.15);
      color: var(--accent);
      margin-bottom: 8px;
      font-weight: 600;
    }}
    .card-title {{
      font-size: 1.1rem;
      font-weight: 700;
      color: #f1f5f9;
      margin-bottom: 6px;
    }}
    .card-math {{
      margin: 12px 0;
      padding: 12px;
      background: rgba(15, 23, 42, 0.6);
      border-radius: 8px;
      font-size: 1.15rem;
      color: #e2e8f0;
      overflow-x: auto;
      text-align: center;
      min-height: 48px;
      display: flex;
      align-items: center;
      justify-content: center;
    }}
    .card-actions {{
      display: flex;
      gap: 8px;
      margin-top: 12px;
    }}
    .copy-btn {{
      flex: 1;
      padding: 8px 12px;
      background: rgba(168, 85, 247, 0.2);
      border: 1px solid var(--purple);
      color: #e9d5ff;
      border-radius: 6px;
      font-size: 0.85rem;
      cursor: pointer;
      font-weight: 600;
      transition: all 0.2s;
    }}
    .copy-btn:hover {{
      background: var(--purple);
      color: #fff;
    }}
    .copy-btn.copied {{
      background: var(--green);
      border-color: var(--green);
      color: #fff;
    }}
    #toast {{
      position: fixed;
      bottom: 24px;
      right: 24px;
      background: #10b981;
      color: #fff;
      padding: 12px 24px;
      border-radius: 8px;
      font-weight: 600;
      box-shadow: 0 4px 12px rgba(0,0,0,0.3);
      display: none;
      z-index: 1000;
    }}
  </style>
  <script>
    window.MathJax = {{
      tex: {{ inlineMath: [['$', '$'], ['\\\\(', '\\\\)']], displayMath: [['$$', '$$']] }},
      svg: {{ fontCache: 'global' }}
    }};
  </script>
  <script src="https://cdn.jsdelivr.net/npm/mathjax@3/es5/tex-svg.js" id="MathJax-script" async></script>
</head>
<body>
  <header>
    <h1>高三理組數學思維本 (2026)<br><span style="font-size: 0.8em; color: #38bdf8;">OMML 數學方程式在線檢索中心</span></h1>
    <p class="lead">收錄全書 20 大專題、共計 {total_slides} 道核心考題之標準參考答案，專為微軟 Word 打造的原生方程式代碼</p>
    <div class="actions">
      <a href="deck2026.html" class="btn btn-primary">📘 返回：高三理組數學思維本(2026) 簡報</a>
      <a href="calculus.html" class="btn btn-sec">📐 微積分專題簡報</a>
      <a href="index.html" class="btn btn-sec">🎯 四校聯考真題簡報</a>
      <a href="高三理組數學思維本_2026_參考答案_OMML對照手冊.docx" class="btn btn-sec" download>📥 下載 Word 對照手冊 (.docx)</a>
      <a href="高三理組數學思維本_2026_參考答案_OMML彙編.xml" class="btn btn-sec" download>📥 下載全書 XML 彙編</a>
    </div>
  </header>

  <div class="filters">
    <input type="text" id="searchBox" class="search-box" placeholder="🔍 搜尋題號、考點關鍵字或公式內容..." />
    <select id="chapterSelect" class="select-box">
      <option value="ALL">全部 20 大專題</option>
    </select>
  </div>

  <div class="container" id="cardsGrid"></div>
  <div id="toast">已複製微軟 Word 原生 OMML 數學方程式代碼！</div>

  <!-- 載入全部 20 個專題資料腳本 -->
"""

for fname in CHAPTER_FILES:
    omml_portal_html += f'  <script src="{fname}"></script>\n'

omml_portal_html += """
  <script>
    document.addEventListener('DOMContentLoaded', () => {
      const DECK = window.DECK || [];
      const grid = document.getElementById('cardsGrid');
      const searchBox = document.getElementById('searchBox');
      const chSelect = document.getElementById('chapterSelect');
      const toast = document.getElementById('toast');

      let allCards = [];

      // 填充下拉選單與收集卡片
      DECK.forEach((chapter, chIdx) => {
        const opt = document.createElement('option');
        opt.value = chIdx;
        opt.textContent = `${chapter.ch || '專題' + (chIdx + 1)} (${(chapter.slides || []).length} 題)`;
        chSelect.appendChild(opt);

        (chapter.slides || []).forEach((slide, sIdx) => {
          allCards.push({
            chapterName: chapter.ch || `專題 ${chIdx + 1}`,
            chapterIdx: chIdx,
            qNum: slide.qNum || `第${sIdx + 1}題`,
            topic: slide.topic || '',
            ans: (slide.solution && slide.solution.ans) ? slide.solution.ans : '',
            omml: (slide.solution && slide.solution.omml) ? slide.solution.omml : ''
          });
        });
      });

      function renderCards(cards) {
        grid.innerHTML = '';
        if (cards.length === 0) {
          grid.innerHTML = '<div style="grid-column: 1/-1; text-align: center; color: var(--sub); padding: 40px;">未找到符合條件的考題答案</div>';
          return;
        }

        cards.forEach((card, idx) => {
          const cardEl = document.createElement('div');
          cardEl.className = 'card';

          let displayMath = card.ans;
          if (displayMath && !displayMath.startsWith('$') && !displayMath.endsWith('$')) {
            displayMath = `$${displayMath}$`;
          }

          cardEl.innerHTML = `
            <div class="card-top">
              <span class="card-tag">${card.chapterName} · ${card.qNum}</span>
              <div class="card-title">${card.topic}</div>
            </div>
            <div class="card-math">${displayMath || '<span style="color:#64748b;">(無公式)</span>'}</div>
            <div class="card-actions">
              <button class="copy-btn" data-idx="${idx}">📋 複製 Word OMML</button>
            </div>
          `;

          const btn = cardEl.querySelector('.copy-btn');
          btn.addEventListener('click', () => {
            if (navigator.clipboard && navigator.clipboard.writeText) {
              navigator.clipboard.writeText(card.omml).then(() => {
                showToast(btn);
              });
            } else {
              const ta = document.createElement('textarea');
              ta.value = card.omml;
              document.body.appendChild(ta);
              ta.select();
              document.execCommand('copy');
              document.body.removeChild(ta);
              showToast(btn);
            }
          });

          grid.appendChild(cardEl);
        });

        if (window.MathJax && MathJax.typesetPromise) {
          MathJax.typesetPromise([grid]);
        }
      }

      function showToast(btn) {
        btn.textContent = '✓ 已複製！';
        btn.classList.add('copied');
        toast.style.display = 'block';
        setTimeout(() => {
          btn.textContent = '📋 複製 Word OMML';
          btn.classList.remove('copied');
          toast.style.display = 'none';
        }, 2200);
      }

      function filterCards() {
        const query = searchBox.value.trim().toLowerCase();
        const chVal = chSelect.value;

        const filtered = allCards.filter(c => {
          const matchCh = (chVal === 'ALL' || c.chapterIdx.toString() === chVal);
          const matchQuery = !query || 
            c.chapterName.toLowerCase().includes(query) ||
            c.qNum.toLowerCase().includes(query) ||
            c.topic.toLowerCase().includes(query) ||
            c.ans.toLowerCase().includes(query);
          return matchCh && matchQuery;
        });

        renderCards(filtered);
      }

      searchBox.addEventListener('input', filterCards);
      chSelect.addEventListener('change', filterCards);

      // 初次渲染
      renderCards(allCards);
    });
  </script>
</body>
</html>
"""

omml_out_demo = os.path.join(DEMO_DIR, "omml.html")
omml_out_assets = os.path.join(ASSETS_DIR, "omml.html")
with open(omml_out_demo, "w", encoding="utf-8") as f:
    f.write(omml_portal_html)
shutil.copyfile(omml_out_demo, omml_out_assets)
print(f"Saved OMML Online Portal to:\n  - {omml_out_demo}\n  - {omml_out_assets}")

# ==============================================================================
# 生成 deck2026.html 主入口簡報
# ==============================================================================
print("\nGenerating Master Presentation Deck: deck2026.html ...")
deck_html_content = f"""<!DOCTYPE html>
<html lang="zh-TW">
<head>
  <meta charset="UTF-8" />
  <meta name="viewport" content="width=device-width, initial-scale=1.0" />
  <title>高三理組數學思維本(2026)</title>
  <meta name="description" content="澳門高三理組數學思維本(2026)全書互動投影教學簡報系統 — 一頁一題、核心知識點、解題步驟、動態圖解與微軟原生 OMML 方程式" />
  <link rel="stylesheet" href="style.css?v=20260924_full" />
  <script>
    window.MathJax = {{
      tex: {{ inlineMath: [['$', '$'], ['\\\\(', '\\\\)']], displayMath: [['$$', '$$']] }},
      svg: {{ fontCache: 'global' }},
      startup: {{ typeset: false }}
    }};
  </script>
  <script src="https://cdn.jsdelivr.net/npm/mathjax@3/es5/tex-svg.js" id="MathJax-script" async></script>
</head>
<body>
  <!-- 封面 -->
  <section id="cover" class="cover">
    <div class="cover-top-bar">
      <button class="cover-teacher-btn" id="coverTeacherBtn" title="教師專用：設定課堂開放進度">⚙️ 課堂進度控制</button>
    </div>
    <div class="cover-inner">
      <p class="cover-eyebrow">澳門中學部數學教材轉化 · 高三高等數學思維探索</p>
      <h1 class="cover-title">高三理組數學思維本(2026)</h1>
      <p class="cover-sub">微積分 · 空間向量 · 立體幾何 · 空間解析幾何 · 基礎概念 · 多項式 · 指對數 · 不等式 · 數列級數 · 解析幾何 · 平面幾何 · 三角函數 · 排列組合概率 · 行列式矩陣 · 複數</p>
      
      <!-- 傳送門導航列 -->
      <div style="margin: 14px 0 18px 0; display:flex; justify-content:center; gap:10px; flex-wrap:wrap;">
        <a href="index.html" class="portal-link-btn" style="display:inline-block;padding:8px 18px;background:rgba(255,255,255,0.08);border:1px solid rgba(255,255,255,0.22);color:#93c5fd;text-decoration:none;border-radius:20px;font-size:0.9rem;font-weight:500;box-shadow:0 2px 8px rgba(0,0,0,0.2);transition:all .2s;">🎯 澳門四校聯考歷屆真題 ›</a>
        <a href="calculus.html" class="portal-link-btn" style="display:inline-block;padding:8px 18px;background:rgba(59,130,246,0.2);border:1px solid rgba(96,165,250,0.4);color:#bfdbfe;text-decoration:none;border-radius:20px;font-size:0.9rem;font-weight:500;box-shadow:0 2px 8px rgba(0,0,0,0.2);transition:all .2s;">📐 微積分專題簡報 ›</a>
        <a href="omml.html" class="portal-link-btn" style="display:inline-block;padding:8px 18px;background:rgba(168,85,247,0.25);border:1px solid rgba(192,132,252,0.4);color:#e9d5ff;text-decoration:none;border-radius:20px;font-size:0.9rem;font-weight:500;box-shadow:0 2px 8px rgba(0,0,0,0.2);transition:all .2s;">📑 OMML 數學方程式庫 ›</a>
      </div>

      <div class="cover-chapters" id="coverChapters"></div>
      <button class="cover-start" id="startBtn">進入高三理組數學思維探索 →</button>
      <p class="cover-hint">用 ← → 鍵換題 · K 鍵展開考點 · A 鍵揭曉解答 · Q 鍵放大題目 · 支援一鍵複製微軟原生 OMML 數學方程式</p>
    </div>
  </section>

  <!-- 主體 -->
  <div id="app" class="app hidden">
    <!-- 頂部列 -->
    <header class="topbar">
      <button class="home-btn" id="homeBtn" title="回封面">✦ 高三理組數學思維本(2026)</button>
      <div class="crumb" id="crumb"></div>
      <div class="top-actions">
        <button class="ghost-btn teacher-ctrl-btn" id="topTeacherBtn" title="課堂開放進度控制">⚙️ 課堂進度控制</button>
        <button class="ghost-btn" id="tocToggle" title="目錄">☰ 目錄</button>
      </div>
    </header>

    <div class="body">
      <!-- 目錄側欄 -->
      <aside class="toc" id="toc"></aside>

      <!-- 舞台 -->
      <main class="stage">
        <div class="slide" id="slide"></div>

        <!-- 底部控制列 -->
        <div class="controlbar">
          <button class="nav-btn" id="prevBtn">‹ 上一頁</button>
          <div class="progress">
            <div class="progress-track"><div class="progress-fill" id="progressFill"></div></div>
            <span class="progress-text" id="progressText">1 / 1</span>
          </div>
          <button class="nav-btn" id="nextBtn">下一頁 ›</button>
        </div>
      </main>
    </div>

    <!-- 授課教具：畫筆畫布、雷射點、浮動工具列 -->
    <canvas id="penCanvas" class="pen-canvas"></canvas>
    <div id="laserDot" class="laser-dot hidden"></div>
    <div id="dock" class="dock">
      <button class="dbtn" id="dkSidebar" title="收合／展開側欄">
        <svg viewBox="0 0 24 24" fill="none" stroke="currentColor" stroke-width="2"><rect x="3" y="4" width="18" height="16" rx="2"/><line x1="9" y1="4" x2="9" y2="20"/></svg>
      </button>
      <button class="dbtn" id="dkPrev" title="上一頁 (←)">
        <svg viewBox="0 0 24 24" fill="none" stroke="currentColor" stroke-width="2.5" stroke-linecap="round" stroke-linejoin="round"><polyline points="15 5 8 12 15 19"/></svg>
      </button>
      <button class="dbtn" id="dkNext" title="下一頁 (→)">
        <svg viewBox="0 0 24 24" fill="none" stroke="currentColor" stroke-width="2.5" stroke-linecap="round" stroke-linejoin="round"><polyline points="9 5 16 12 9 19"/></svg>
      </button>
      <div class="dsep"></div>
      <button class="dbtn" id="dkLaser" title="雷射筆">
        <svg viewBox="0 0 24 24" fill="none" stroke="currentColor" stroke-width="2"><circle cx="12" cy="12" r="7"/><circle cx="12" cy="12" r="2.5" fill="currentColor" stroke="none"/></svg>
      </button>
      <button class="dbtn" id="dkPen" title="畫筆">
        <svg viewBox="0 0 24 24" fill="none" stroke="currentColor" stroke-width="2" stroke-linecap="round" stroke-linejoin="round"><path d="M12 20h9"/><path d="M16.5 3.5a2.12 2.12 0 0 1 3 3L7 19l-4 1 1-4Z"/></svg>
      </button>
      <div class="dcolors hidden" id="dkColors">
        <button class="dcolor active" data-c="#e11d48" style="--sc:#e11d48" title="紅"></button>
        <button class="dcolor" data-c="#2563eb" style="--sc:#2563eb" title="藍"></button>
        <button class="dcolor" data-c="#059669" style="--sc:#059669" title="綠"></button>
        <button class="dcolor" data-c="#f59e0b" style="--sc:#f59e0b" title="黃"></button>
        <button class="dcolor" data-c="#111827" style="--sc:#111827" title="黑"></button>
        <button class="dbtn derase" id="dkErase" title="橡皮擦">
          <svg viewBox="0 0 24 24" fill="none" stroke="currentColor" stroke-width="2" stroke-linecap="round" stroke-linejoin="round"><path d="M4 15l7-7 6 6-4 4H8z"/><line x1="4" y1="21" x2="20" y2="21"/></svg>
        </button>
      </div>
      <div class="dsep"></div>
      <button class="dbtn" id="dkKnowledge" title="核心知識點與公式 (K)">
        <svg viewBox="0 0 24 24" fill="none" stroke="currentColor" stroke-width="2"><path d="M2 3h6a4 4 0 0 1 4 4v14a3 3 0 0 0-3-3H2z"/><path d="M22 3h-6a4 4 0 0 0-4 4v14a3 3 0 0 1 3-3h7z"/></svg>
      </button>
      <button class="dbtn" id="dkAns" title="顯示／隱藏答案 (A)">
        <svg viewBox="0 0 24 24" fill="none" stroke="currentColor" stroke-width="2" stroke-linecap="round" stroke-linejoin="round"><circle cx="12" cy="12" r="10"/><path d="m9 12 2 2 4-4"/></svg>
      </button>
      <button class="dbtn" id="dkZoom" title="放大題目與圖解 (Q)">
        <svg viewBox="0 0 24 24" fill="none" stroke="currentColor" stroke-width="2" stroke-linecap="round" stroke-linejoin="round"><circle cx="11" cy="11" r="8"/><line x1="21" y1="21" x2="16.65" y2="16.65"/><line x1="11" y1="8" x2="11" y2="14"/><line x1="8" y1="11" x2="14" y2="11"/></svg>
      </button>
      <button class="dbtn" id="dkFullscreen" title="全螢幕 (F)">
        <svg viewBox="0 0 24 24" fill="none" stroke="currentColor" stroke-width="2"><path d="M8 3H5a2 2 0 0 0-2 2v3m18 0V5a2 2 0 0 0-2-2h-3m0 18h3a2 2 0 0 0 2-2v-3M3 16v3a2 2 0 0 0 2 2h3"/></svg>
      </button>
    </div>
  </div>

  <!-- 放大模態窗 -->
  <div id="zoomModal" class="zoom-modal hidden">
    <div class="zoom-modal-backdrop" id="zoomBackdrop"></div>
    <div class="zoom-modal-dialog">
      <button class="zoom-modal-close" id="zoomClose" title="關閉 (Esc)">✕</button>
      <div class="zoom-modal-body" id="zoomBody"></div>
    </div>
  </div>

  <!-- 核心知識點與必背公式放大模態窗 -->
  <div id="knowledgeModal" class="zoom-modal hidden">
    <div class="zoom-modal-backdrop" id="knowledgeBackdrop"></div>
    <div class="zoom-modal-dialog">
      <button class="zoom-modal-close" id="knowledgeClose" title="關閉 (Esc)">✕</button>
      <div class="zoom-modal-body" id="knowledgeBody"></div>
    </div>
  </div>

  <!-- 教師控制面板模態窗 -->
  <div id="teacherModal" class="teacher-modal hidden">
    <div class="teacher-modal-backdrop" id="teacherBackdrop"></div>
    <div class="teacher-modal-dialog">
      <div class="teacher-modal-header">
        <h3 class="teacher-modal-title">⚙️ 課堂實驗開放進度控制</h3>
        <button class="teacher-modal-close" id="teacherClose" title="關閉">✕</button>
      </div>
      <div class="teacher-modal-body">
        <div id="teacherPinView" class="teacher-pin-view">
          <p class="pin-hint">請輸入教師 PIN 碼（預設：8888）</p>
          <div class="pin-inputs">
            <input type="password" id="pinInput" maxlength="6" class="pin-field" placeholder="****" autofocus />
            <button id="pinSubmitBtn" class="pin-btn">解鎖控制台</button>
          </div>
          <p id="pinError" class="pin-error hidden">PIN 碼錯誤，請重新輸入</p>
        </div>
        <div id="teacherControlView" class="teacher-control-view hidden">
          <div class="preset-buttons">
            <span class="preset-label">快速進度：</span>
            <button class="preset-btn" data-preset="w1-2">第 1~2 週 (專題 1~4)</button>
            <button class="preset-btn" data-preset="w3-4">第 3~4 週 (專題 5~8)</button>
            <button class="preset-btn" data-preset="w5-6">第 5~6 週 (專題 9~14)</button>
            <button class="preset-btn" data-preset="all">全部開放 (專題 1~20)</button>
          </div>
          <div class="chapter-check-list" id="chapterCheckList"></div>
          <div class="teacher-actions">
            <button id="saveScopeBtn" class="save-scope-btn">💾 儲存並在當前生效</button>
            <button id="copyStudentUrlBtn" class="copy-url-btn">📋 複製學生網址</button>
            <button id="showQrBtn" class="qr-btn">📱 大螢幕 QR 碼</button>
          </div>
        </div>
      </div>
    </div>
  </div>

  <!-- QR Code 投影彈窗 -->
  <div id="qrModal" class="zoom-modal hidden">
    <div class="zoom-modal-backdrop" id="qrBackdrop"></div>
    <div class="zoom-modal-dialog qr-dialog">
      <button class="zoom-modal-close" id="qrClose" title="關閉">✕</button>
      <div class="qr-body">
        <h3 class="qr-title">📱 學生請掃描二維碼同步課堂進度</h3>
        <p class="qr-subtitle" id="qrSubtitle">僅開放當前勾選之專題</p>
        <div id="qrcodeBox" class="qrcode-box"></div>
        <p class="qr-hint">學生掃描後，iPad 將自動鎖定在指定授課章節</p>
      </div>
    </div>
  </div>

  <!-- QR Code 函式庫 -->
  <script src="qrcode.min.js"></script>

  <!-- 向量與動態圖解引擎 -->
  <script src="svg.js"></script>

  <!-- 全書 20 個專題資料腳本 -->
"""

for fname in CHAPTER_FILES:
    deck_html_content += f'  <script src="{fname}"></script>\n'

deck_html_content += """
  <!-- 核心渲染與互動引擎 -->
  <script src="engine.js?v=20260924_full"></script>
</body>
</html>
"""

deck_out_demo = os.path.join(DEMO_DIR, "deck2026.html")
deck_out_assets = os.path.join(ASSETS_DIR, "deck2026.html")
with open(deck_out_demo, "w", encoding="utf-8") as f:
    f.write(deck_html_content)
shutil.copyfile(deck_out_demo, deck_out_assets)
print(f"Saved Master Deck to:\n  - {deck_out_demo}\n  - {deck_out_assets}")

# ==============================================================================
# 更新 index.html 與 calculus.html 傳送門按鈕
# ==============================================================================
print("\nUpdating portal links in index.html & calculus.html ...")

def update_portal_in_file(file_path):
    with open(file_path, "r", encoding="utf-8") as f:
        html = f.read()
    
    # Check if deck2026 link already exists
    if "deck2026.html" not in html:
        # insert deck2026 portal link button
        btn_deck2026 = '<a href="deck2026.html" class="portal-link-btn" style="display:inline-block;padding:8px 18px;background:rgba(56,189,248,0.25);border:1px solid rgba(56,189,248,0.5);color:#7dd3fc;text-decoration:none;border-radius:20px;font-size:0.9rem;font-weight:600;box-shadow:0 2px 8px rgba(0,0,0,0.2);transition:all .2s;">📘 高三理組數學思維本(2026) 全書完整版 ›</a>'
        html = re.sub(r'(<div style="margin:\s*14px 0 18px 0;[^>]*>)', r'\1\n        ' + btn_deck2026, html, count=1)
        with open(file_path, "w", encoding="utf-8") as f:
            f.write(html)
        print(f"  Added deck2026 link to {file_path}")

update_portal_in_file(os.path.join(DEMO_DIR, "index.html"))
update_portal_in_file(os.path.join(ASSETS_DIR, "index.html"))
update_portal_in_file(os.path.join(DEMO_DIR, "calculus.html"))
update_portal_in_file(os.path.join(ASSETS_DIR, "calculus.html"))

print("\n=== All builds and deployments completed successfully! ===")
